import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
# Set slide dimensions to widescreen 16:9
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Dark modern theme palette
BG_DARK = RGBColor(15, 23, 42)      # #0f172a
ACCENT_ROSE = RGBColor(244, 63, 94)  # #f43f5e
ACCENT_AMBER = RGBColor(245, 158, 11) # #f59e0b
TEXT_WHITE = RGBColor(255, 255, 255)
TEXT_MUTED = RGBColor(148, 163, 184) # #94a3b8

slides_data = [
    {
        "title": "OmniDine AI",
        "subtitle": "Interconnected Smart SaaS Operating System for Restaurants\n\nTeam Name: Crafterstudio  |  Team Lead: Bhavit Jangra\nVibeAthon 6.0 Hackathon Submission",
        "type": "title"
    },
    {
        "title": "The Problem in Restaurant Operations",
        "points": [
            "Disconnected Stakeholders: Customer, Waiter, Kitchen Chef, and Manager operate in isolated silos.",
            "Paper Ticket Errors: Manual paper chits cause order mistakes, delays, and lost revenues.",
            "Lack of Real-Time SLA Auditing: Managers lack live visibility into kitchen prep velocity and floor occupancy."
        ]
    },
    {
        "title": "The Solution — OmniDine AI Platform",
        "points": [
            "100% Interconnected 4-Role Pipeline: Real-time synchronization across Customer, Waiter, Chef KDS, and Manager Admin.",
            "Gourmet Pure Veg Dining: Authentic Indian menu with Thalis, Combos, and customizable spice levels.",
            "Digital KDS & Instant UPI Checkout: Live order ticket timers, scannable UPI QR code payment, and PDF receipts."
        ]
    },
    {
        "title": "Gourmet Pure Veg Menu & UPI Payments",
        "points": [
            "Signature Indian Dishes: Paneer Butter Masala, Dal Makhani Royale, Hyderabadi Biryani, Amritsari Chole Kulcha.",
            "Special Combos & Buffets: Maharaja Unlimited Veg Buffet, Shahi Punjabi Thali, South Indian Feast Thali.",
            "Seamless Digital Payments: Integrated GPay, PhonePe, Paytm QR code payment and split bill management."
        ]
    },
    {
        "title": "Tech Stack & Architecture",
        "points": [
            "Frontend Stack: React 18, Vite, Tailwind CSS, Lucide Icons, Recharts Analytics.",
            "State Engine & Resilience: Context API state engine with React Error Boundary crash isolation.",
            "Multi-Cloud Ready: Vercel SPA rewritten routes (vercel.json) & Netlify configuration (_redirects)."
        ]
    },
    {
        "title": "Live Product Walkthrough",
        "points": [
            "1. Custom Onboarding: Configure restaurant name, logo image/emoji, and currency ($/₹).",
            "2. Active Order Banner: Glowing green banner displays live status at the top of customer menu.",
            "3. KDS Prep Bumping: Chefs bump order stage from Received ➔ In Prep ➔ Ready ➔ Served in real time.",
            "4. Manager Audit Log: Operational analytics dashboard updates sales revenue dynamically."
        ]
    },
    {
        "title": "Future Roadmap & Beta Features",
        "points": [
            "[BETA] Multi-Branch Synchronization: Cross-location franchise sales aggregation.",
            "[BETA] Dynamic Pricing Engine: AI demand surge auto-adjustments during peak footfall.",
            "[BETA] Automated Supplier PO Generation: Direct REST API purchase orders triggered by low stock."
        ]
    },
    {
        "title": "Summary & Submission Details",
        "points": [
            "GitHub Repository: https://github.com/CrafterStudioIndia/Omni-Dinner-App-",
            "Team Lead Name: Bhavit Jangra (Crafterstudio)",
            "Live Web App: http://localhost:3000/",
            "Thank You! Built for VibeAthon 6.0 Hackathon."
        ]
    }
]

blank_slide_layout = prs.slide_layouts[6]

for slide_info in slides_data:
    slide = prs.slides.add_slide(blank_slide_layout)
    
    # Background fill
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_DARK
    
    # Slide Title
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.2))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = slide_info["title"]
    p_title.font.size = Pt(40)
    p_title.font.bold = True
    p_title.font.color.rgb = ACCENT_ROSE if slide_info.get("type") == "title" else TEXT_WHITE

    if slide_info.get("type") == "title":
        p_title.alignment = PP_ALIGN.CENTER
        title_box.top = Inches(2.2)
        
        sub_box = slide.shapes.add_textbox(Inches(1.5), Inches(3.8), Inches(10.333), Inches(2.5))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = slide_info["subtitle"]
        p_sub.font.size = Pt(22)
        p_sub.font.color.rgb = TEXT_MUTED
        p_sub.alignment = PP_ALIGN.CENTER
    else:
        # Bullet points container
        content_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(4.5))
        tf_content = content_box.text_frame
        tf_content.word_wrap = True
        
        for idx, pt_text in enumerate(slide_info["points"]):
            p = tf_content.add_paragraph() if idx > 0 else tf_content.paragraphs[0]
            p.text = f"• {pt_text}"
            p.font.size = Pt(22)
            p.font.color.rgb = TEXT_MUTED
            p.space_after = Pt(20)

output_path = r"c:\Users\SHAKTI\OneDrive\Documents\OmniDine_Crafterstudio_Presentation.pptx"
prs.save(output_path)
print(f"Presentation saved successfully to {output_path}")
